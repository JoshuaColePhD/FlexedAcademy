"""Upload-boundary tests for reusable lesson-plan templates."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend import template_intake
from backend.errors import AppError


def _structure(*, total_text_chars: int = 300, cells: list[str] | None = None) -> dict:
    return {
        "total_text_chars": total_text_chars,
        "tables": [{"sample_rows": [cells or ["Learning target", "Monday", "Tuesday"]]}],
    }


def _use_structure(monkeypatch, structure: dict) -> None:
    monkeypatch.setattr(template_intake, "_extract_docx_structure", lambda _path: structure)
    monkeypatch.setattr(template_intake, "_extract_pdf_structure", lambda _path: structure)


def test_blank_template_with_labels_is_accepted(monkeypatch):
    _use_structure(monkeypatch, _structure(cells=["Learning target", "Vocabulary", "I Do / We Do / You Do"]))

    template_intake.require_blank_template(Path("blank.docx"), ".docx")


def test_completed_template_is_rejected(monkeypatch):
    _use_structure(monkeypatch, _structure(total_text_chars=9_000))

    with pytest.raises(AppError, match="completed lesson plan") as error:
        template_intake.require_blank_template(Path("completed.docx"), ".docx")

    assert error.value.code == "template_not_blank"


def test_template_with_personal_contact_data_is_rejected(monkeypatch):
    _use_structure(monkeypatch, _structure(cells=["Teacher: Jane Smith", "jane.smith@example.org"]))

    with pytest.raises(AppError, match="email address") as error:
        template_intake.require_blank_template(Path("contact.docx"), ".docx")

    assert error.value.code == "template_contains_personal_data"


def test_docx_zip_with_unsafe_expanded_size_is_rejected(tmp_path):
    path = tmp_path / "inflate.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (50 * 1024 * 1024 + 1))

    with pytest.raises(AppError) as error:
        template_intake.check_docx_zip(path)

    assert error.value.code == "file_too_large"


def test_docx_zip_with_unsafe_compression_ratio_is_rejected(tmp_path):
    path = tmp_path / "ratio.docx"
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (1024 * 1024))

    with pytest.raises(AppError) as error:
        template_intake.check_docx_zip(path)

    assert error.value.code == "file_too_large"


def test_docx_structure_extracts_cell_visual_design(tmp_path):
    """The template contract includes visual intent, not only cell text."""
    path = tmp_path / "visual-template.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cell = table.cell(0, 0)
    cell.text = "Learning target"
    cell.width = Inches(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run = cell.paragraphs[0].runs[0]
    run.font.name = "Aptos Display"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    run.bold = True
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "EADCF8")
    tc_pr.append(shading)
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), "FF0000")
    bottom.set(qn("w:sz"), "8")
    borders.append(bottom)
    tc_pr.append(borders)
    table.cell(0, 1).text = "Monday"
    table.cell(1, 0).text = "Objective"
    doc.save(path)

    structure = template_intake._extract_docx_structure(path)
    visual = structure["tables"][0]["visual_rows"][0][0]
    assert visual["fill_hex"] == "EADCF8"
    assert visual["width_dxa"]
    assert visual["vertical_align"] == "center"
    assert visual["font_colors"] == ["123456"]
    assert visual["font_sizes_pt"] == [14.0]
    assert visual["bold"] is True
    assert visual["borders"]["bottom"]["color"] == "FF0000"
